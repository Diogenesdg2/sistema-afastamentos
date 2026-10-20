import os
import platform
import subprocess
from tkinter import messagebox
import sqlite3

# --- Importações do PDF (ReportLab) ---
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from database import conectar_db, calcular_dias_beneficio_aquisitivo

# --- Importações do Excel (OpenPyXL) ---
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# =========================================================================
# ========================= EXPORTAÇÃO EM PDF =============================
# =========================================================================

def gerar_pdf_afastamentos(ids_lista, caminho_arquivo):
    if not ids_lista:
        messagebox.showwarning("Aviso", "Não há dados na tabela para gerar o relatório.")
        return

    doc = SimpleDocTemplate(caminho_arquivo, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    elementos = []
    estilos = getSampleStyleSheet()

    estilo_titulo = ParagraphStyle('TituloPrincipal', parent=estilos['Heading1'], fontSize=16, textColor=colors.HexColor("#1D4ED8"), alignment=1, spaceAfter=20)
    estilo_sub = ParagraphStyle('SubTitulo', parent=estilos['Heading2'], fontSize=11, textColor=colors.white, backColor=colors.HexColor("#3B82F6"), spaceBefore=15, spaceAfter=5, leftIndent=5, rightIndent=5, borderPadding=6)
    estilo_secao = ParagraphStyle('SecaoInterna', parent=estilos['Normal'], fontSize=9, fontName="Helvetica-Bold", textColor=colors.HexColor("#1E293B"), spaceBefore=6, spaceAfter=3)
    estilo_texto = estilos['Normal']

    elementos.append(Paragraph("Relatório Detalhado de Afastamentos INSS", estilo_titulo))

    conn = conectar_db()
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    placeholders = ','.join('?' * len(ids_lista))
    
    query = f"SELECT * FROM afastamentos WHERE id IN ({placeholders}) ORDER BY empresa, nome_empregado"
    cursor.execute(query, ids_lista)
    registros = cursor.fetchall()

    for linha in registros:
        bloco_funcionario = []
        
        id_afastamento = linha['id']
        empresa = linha['empresa'] or '-'
        cod_empresa = linha['cod_empresa'] if 'cod_empresa' in linha.keys() and linha['cod_empresa'] else '-'
        nome = linha['nome_empregado'] or '-'
        cod_func = linha['cod_empregado'] or '-'
        cpf_func = linha['cpf'] if 'cpf' in linha.keys() and linha['cpf'] else '-'
        motivo = linha['tipo_afastamento'] or '-'
        ult_dia = linha['ultimo_dia_trab'] or '-'
        obrigacao = linha['data_obrigacao'] or '-'
        retornou = "Sim" if linha['retornou'] else "Não (Afastado)"
        obs = linha['observacoes'] if 'observacoes' in linha.keys() and str(linha['observacoes']).strip() != "" else "-"

        aposentado_str = "Sim" if str(linha['aposentado']) in ("1", "Sim", "True") else "Não"
        domestica_str = "Sim" if 'domestica' in linha.keys() and str(linha['domestica']) in ("1", "Sim", "True") else "Não"

        # CÁLCULO DE FÉRIAS PARA O RELATÓRIO
        status_ferias = "-"
        if 'controlar_ferias' in linha.keys() and linha['controlar_ferias'] == 1 and linha['per_aquisitivo_de']:
            dias_ferias, qtd_reg = calcular_dias_beneficio_aquisitivo(cpf_func, cod_empresa, cod_func, linha['per_aquisitivo_de'], linha['per_aquisitivo_ate'], id_afastamento)
            status_ferias = f"{dias_ferias} / 180 dias concedidos"
            if dias_ferias >= 180: status_ferias += " (DIREITO PERDIDO)"

        tags_func = f" | 👴 Aposentado: {aposentado_str} | 🧹 Doméstica: {domestica_str}"
        bloco_funcionario.append(Paragraph(f"<b>🏢 {empresa}</b> (Cód: {cod_empresa}) | 👤 {nome} (Cód: {cod_func}) | CPF: {cpf_func}{tags_func}", estilo_sub))

        dados_basicos = [
            [Paragraph("<b>Motivo Afastamento:</b>", estilo_texto), Paragraph(motivo, estilo_texto),
             Paragraph("<b>Status de Férias:</b>", estilo_texto), Paragraph(status_ferias, estilo_texto)],
             
            [Paragraph("<b>Data Obrigação:</b>", estilo_texto), Paragraph(obrigacao, estilo_texto),
             Paragraph("<b>Já Retornou?:</b>", estilo_texto), Paragraph(retornou, estilo_texto)]
        ]
        
        tabela_basica = Table(dados_basicos, colWidths=[120, 150, 110, 155])
        tabela_basica.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (0,-1), colors.HexColor("#F3F4F6")),
            ('BACKGROUND', (2,0), (2,-1), colors.HexColor("#F3F4F6")),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.lightgrey),
            ('BOX', (0,0), (-1,-1), 0.5, colors.lightgrey),
            ('BOTTOMPADDING', (0,0), (-1,-1), 5),
            ('TOPPADDING', (0,0), (-1,-1), 5),
        ]))
        bloco_funcionario.append(tabela_basica)
        bloco_funcionario.append(Spacer(1, 8))

        cursor.execute("SELECT codigo, descricao FROM cids_historico WHERE id_afastamento=? ORDER BY id ASC", (id_afastamento,))
        cids = cursor.fetchall()
        
        bloco_funcionario.append(Paragraph("<b>Diagnósticos Médicos (CIDs Registrados):</b>", estilo_secao))
        dados_cids = [
            [Paragraph("<b>Cód. CID</b>", estilo_texto), Paragraph("<b>Descrição / Patologia</b>", estilo_texto)]
        ]
        
        if cids:
            for c in cids:
                dados_cids.append([
                    Paragraph(f"<b>{c['codigo']}</b>", estilo_texto),
                    Paragraph(c['descricao'] or "Descrição não informada", estilo_texto)
                ])
        else:
            cid_legado = linha['cid'] or '-'
            dados_cids.append([
                Paragraph(f"<b>{cid_legado}</b>", estilo_texto),
                Paragraph("Diagnóstico inicial registrado", estilo_texto)
            ])

        tabela_cid = Table(dados_cids, colWidths=[90, 445])
        tabela_cid.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#E0E7FF")),  
            ('ALIGN', (0,0), (0,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.lightgrey),
            ('BOX', (0,0), (-1,-1), 0.8, colors.HexColor("#6366F1")),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ('TOPPADDING', (0,0), (-1,-1), 4),
        ]))
        bloco_funcionario.append(tabela_cid)
        bloco_funcionario.append(Spacer(1, 8))

        cursor.execute("SELECT id, data_analise, gerou_pericia, numero_beneficio, situacao, concessao_de, concessao_ate FROM analises_inss WHERE id_afastamento=? ORDER BY id ASC", (id_afastamento,))
        analises = cursor.fetchall()
        
        cursor.execute("SELECT data_agendamento, data_atendimento, especie, numero_beneficio, situacao, concessao_de, concessao_ate FROM pericias_inss WHERE id_afastamento=? AND (id_analise_mae IS NULL OR id_analise_mae = '') ORDER BY id ASC", (id_afastamento,))
        pericias_orfas = cursor.fetchall()
        
        if analises or pericias_orfas:
            bloco_funcionario.append(Paragraph("<b>Trâmites INSS (Análises e Perícias Vinculadas):</b>", estilo_secao))
            dados_inss = [
                [Paragraph("<b>Trâmite</b>", estilo_texto), Paragraph("<b>Data / Agend.</b>", estilo_texto), Paragraph("<b>Perícia? / Esp.</b>", estilo_texto), Paragraph("<b>Nº Benefício</b>", estilo_texto), Paragraph("<b>Situação</b>", estilo_texto), Paragraph("<b>Período (De - Até)</b>", estilo_texto)]
            ]
            
            estilo_tabela_inss = [
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#E5E7EB")),
                ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                ('INNERGRID', (0,0), (-1,-1), 0.5, colors.lightgrey),
                ('BOX', (0,0), (-1,-1), 1, colors.grey),
                ('BOTTOMPADDING', (0,0), (-1,-1), 4),
                ('TOPPADDING', (0,0), (-1,-1), 4),
            ]
            
            row_idx = 1
            for i, an in enumerate(analises, 1):
                id_an = an['id']
                de_ate = f"{an['concessao_de'] or '-'} a {an['concessao_ate'] or '-'}" if (an['concessao_de'] or an['concessao_ate']) else "-"
                
                dados_inss.append([
                    Paragraph(f"<b>{i}ª Análise</b>", estilo_texto),
                    Paragraph(an['data_analise'] or '-', estilo_texto),
                    Paragraph(an['gerou_pericia'] or '-', estilo_texto),
                    Paragraph(an['numero_beneficio'] or '-', estilo_texto),
                    Paragraph(an['situacao'] or '-', estilo_texto),
                    Paragraph(de_ate, estilo_texto)
                ])
                estilo_tabela_inss.append(('BACKGROUND', (0, row_idx), (-1, row_idx), colors.HexColor("#DBEAFE"))) 
                row_idx += 1
                
                cursor.execute("SELECT data_agendamento, data_atendimento, especie, numero_beneficio, situacao, concessao_de, concessao_ate FROM pericias_inss WHERE id_analise_mae=? ORDER BY id ASC", (id_an,))
                pericias = cursor.fetchall()
                
                for j, per in enumerate(pericias, 1):
                    p_de_ate = f"{per['concessao_de'] or '-'} a {per['concessao_ate'] or '-'}" if (per['concessao_de'] or per['concessao_ate']) else "-"
                    dados_inss.append([
                        Paragraph(f"<i> ↳ Perícia {j}</i>", estilo_texto),
                        Paragraph(f"Ag: {per['data_agendamento'] or '-'}<br/>At: {per['data_atendimento'] or '-'}", estilo_texto),
                        Paragraph(f"Esp: {per['especie'] or '-'}", estilo_texto),
                        Paragraph(per['numero_beneficio'] or '-', estilo_texto),
                        Paragraph(per['situacao'] or '-', estilo_texto),
                        Paragraph(p_de_ate, estilo_texto)
                    ])
                    estilo_tabela_inss.append(('BACKGROUND', (0, row_idx), (-1, row_idx), colors.HexColor("#FEF3C7"))) 
                    row_idx += 1
            
            for k, per in enumerate(pericias_orfas, 1):
                p_de_ate = f"{per['concessao_de'] or '-'} a {per['concessao_ate'] or '-'}" if (per['concessao_de'] or per['concessao_ate']) else "-"
                dados_inss.append([
                    Paragraph(f"<i> ↳ Perícia Sem Vínculo {k}</i>", estilo_texto),
                    Paragraph(f"Ag: {per['data_agendamento'] or '-'}<br/>At: {per['data_atendimento'] or '-'}", estilo_texto),
                    Paragraph(f"Esp: {per['especie'] or '-'}", estilo_texto),
                    Paragraph(per['numero_beneficio'] or '-', estilo_texto),
                    Paragraph(per['situacao'] or '-', estilo_texto),
                    Paragraph(p_de_ate, estilo_texto)
                ])
                estilo_tabela_inss.append(('BACKGROUND', (0, row_idx), (-1, row_idx), colors.HexColor("#FEF3C7")))
                row_idx += 1
            
            tabela_inss = Table(dados_inss, colWidths=[75, 85, 80, 95, 80, 120])
            tabela_inss.setStyle(TableStyle(estilo_tabela_inss))
            bloco_funcionario.append(tabela_inss)
            bloco_funcionario.append(Spacer(1, 8))

        bloco_funcionario.append(Paragraph("<b>Tabela de Atestados Sucessivos (Prorrogações):</b>", estilo_secao))
        dados_atestados = [["Ordem", "Início (De)", "Final (Até)", "Dias", "Observações"]]
        
        cursor.execute("SELECT data_de, data_ate, qtd_dias, observacoes FROM atestados_historico WHERE id_afastamento=? ORDER BY id ASC", (id_afastamento,))
        historico = cursor.fetchall()
        if historico:
            for i, hist in enumerate(historico, start=1):
                dados_atestados.append([f"{i}º Atestado", hist['data_de'] or '-', hist['data_ate'] or '-', str(hist['qtd_dias']) if hist['qtd_dias'] else '-', Paragraph(hist['observacoes'] or '-', estilo_texto)])
        else:
            dados_atestados.append(["-", "-", "-", "-", Paragraph("Nenhum atestado registrado.", estilo_texto)])

        tabela_hist = Table(dados_atestados, colWidths=[70, 75, 75, 45, 270])
        tabela_hist.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#E5E7EB")),
            ('TEXTCOLOR', (0,0), (-1,0), colors.black),
            ('ALIGN', (0,0), (-1,0), 'CENTER'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('ALIGN', (0,1), (3,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.lightgrey),
            ('BOX', (0,0), (-1,-1), 1, colors.grey),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ('TOPPADDING', (0,0), (-1,-1), 4),
        ]))
        bloco_funcionario.append(tabela_hist)

        if obs != "-":
            bloco_funcionario.append(Spacer(1, 4))
            bloco_funcionario.append(Paragraph(f"<b>Obs Gerais do Processo:</b> {obs}", estilo_texto))

        bloco_funcionario.append(Spacer(1, 16))
        elementos.append(KeepTogether(bloco_funcionario))

    conn.close()

    try:
        doc.build(elementos)
        if platform.system() == 'Windows': os.startfile(caminho_arquivo)
        elif platform.system() == 'Darwin': subprocess.call(('open', caminho_arquivo))
        else: subprocess.call(('xdg-open', caminho_arquivo))
    except Exception as e:
        messagebox.showerror("Erro", f"Erro ao gerar o PDF:\n{str(e)}")


def gerar_pdf_gestantes(ids_lista, caminho_arquivo):
    if not ids_lista:
        messagebox.showwarning("Aviso", "Não há dados na tabela para gerar o relatório.")
        return

    doc = SimpleDocTemplate(caminho_arquivo, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    elementos = []
    estilos = getSampleStyleSheet()

    estilo_titulo = ParagraphStyle('TituloPrincipal', parent=estilos['Heading1'], fontSize=16, textColor=colors.HexColor("#DB2777"), alignment=1, spaceAfter=20)
    estilo_sub = ParagraphStyle('SubTitulo', parent=estilos['Heading2'], fontSize=11, textColor=colors.white, backColor=colors.HexColor("#BE185D"), spaceBefore=15, spaceAfter=5, leftIndent=5, rightIndent=5, borderPadding=6)
    estilo_texto = estilos['Normal']

    elementos.append(Paragraph("Relatório de Acompanhamento de Gestantes", estilo_titulo))

    conn = conectar_db()
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    placeholders = ','.join('?' * len(ids_lista))
    
    query = f"SELECT * FROM gestantes WHERE id IN ({placeholders}) ORDER BY empresa, nome_empregado"
    cursor.execute(query, ids_lista)
    registros = cursor.fetchall()

    for linha in registros:
        bloco = []
        
        empresa = linha['empresa'] or '-'
        cod_empresa = linha['cod_empresa'] if 'cod_empresa' in linha.keys() and linha['cod_empresa'] else '-'
        nome = linha['nome_empregado'] or '-'
        cod_func = linha['cod_empregado'] if 'cod_empregado' in linha.keys() and linha['cod_empregado'] else '-'
        status = linha['status'] or '-'
        dt_exame = linha['data_exame'] or '-'
        semanas = str(linha['semanas_exame']) if linha['semanas_exame'] else '-'
        dt_40 = linha['data_37_semanas'] or '-'
        estab = linha['estabilidade'] or '-'
        
        bloco.append(Paragraph(f"<b>🏢 {empresa}</b> (Cód: {cod_empresa}) | 🤰 {nome} (Cód: {cod_func})", estilo_sub))

        dados_basicos = [
            [Paragraph("<b>Status Atual:</b>", estilo_texto), Paragraph(status, estilo_texto),
             Paragraph("<b>Data do Exame:</b>", estilo_texto), Paragraph(dt_exame, estilo_texto)],
             
            [Paragraph("<b>Semanas Constatadas:</b>", estilo_texto), Paragraph(semanas, estilo_texto),
             Paragraph("<b>Projeção 40ª Sem. (Parto):</b>", estilo_texto), Paragraph(dt_40, estilo_texto)],

            [Paragraph("<b>Fim da Estabilidade:</b>", estilo_texto), Paragraph(estab, estilo_texto),
             Paragraph("", estilo_texto), Paragraph("", estilo_texto)]
        ]
        
        tabela_basica = Table(dados_basicos, colWidths=[130, 140, 130, 135])
        tabela_basica.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#FDF2F8")),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.lightgrey),
            ('BOX', (0,0), (-1,-1), 0.5, colors.lightgrey),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ('TOPPADDING', (0,0), (-1,-1), 6),
        ]))
        bloco.append(tabela_basica)
        bloco.append(Spacer(1, 15))

        elementos.append(KeepTogether(bloco))

    conn.close()

    try:
        doc.build(elementos)
        if platform.system() == 'Windows': os.startfile(caminho_arquivo)
        elif platform.system() == 'Darwin': subprocess.call(('open', caminho_arquivo))
        else: subprocess.call(('xdg-open', caminho_arquivo))
    except Exception as e:
        messagebox.showerror("Erro", f"Erro ao gerar o PDF:\n{str(e)}")


# =========================================================================
# ========================= EXPORTAÇÃO EM EXCEL ===========================
# =========================================================================

def auto_ajustar_colunas(ws):
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    linhas_texto = str(cell.value).split('\n')
                    for linha in linhas_texto:
                        if len(linha) > max_length:
                            max_length = len(linha)
            except:
                pass
        
        adjusted_width = min(max_length + 2, 55) 
        ws.column_dimensions[column].width = adjusted_width

def gerar_excel_afastamentos(ids_lista, caminho_arquivo):
    if not HAS_OPENPYXL:
        messagebox.showerror("Biblioteca Ausente", "Para gerar relatórios em Excel, é necessário instalar a biblioteca openpyxl.\n\nAbra o terminal do seu projeto e digite:\n\npip install openpyxl")
        return
        
    if not ids_lista:
        messagebox.showwarning("Aviso", "Não há dados na tabela para exportar.")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Afastamentos INSS"

    preenchimento_cabecalho = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid") 
    fonte_cabecalho = Font(color="FFFFFF", bold=True)
    
    alinhamento_centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    alinhamento_esq_topo = Alignment(horizontal="left", vertical="top", wrap_text=True)
    
    borda_fina = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    cabecalhos = [
        "ID", "Empresa", "Cód Emp.", "Funcionário", "CPF", "Cód Func.", "Motivo", 
        "Aposentado/Doméstica", "Últ. Dia Trab.", "Data Obrigação", "Período De", 
        "Período Até", "Dias Totais", "Retornou?", "Controle Férias (Art 133)",
        "Diagnósticos Médicos (CIDs)", "Trâmites INSS (Análises e Perícias)", 
        "Histórico de Atestados", "Observações Gerais"
    ]
    
    ws.append(cabecalhos)

    for cell in ws[1]:
        cell.fill = preenchimento_cabecalho
        cell.font = fonte_cabecalho
        cell.alignment = alinhamento_centro
        cell.border = borda_fina

    conn = conectar_db()
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    placeholders = ','.join('?' * len(ids_lista))
    
    query = f"SELECT * FROM afastamentos WHERE id IN ({placeholders}) ORDER BY empresa, nome_empregado"
    cursor.execute(query, ids_lista)
    registros = cursor.fetchall()

    for linha in registros:
        cpf_func = linha['cpf'] if 'cpf' in linha.keys() else ""
        cod_func = linha['cod_empregado'] or ""
        cod_emp = linha['cod_empresa'] if 'cod_empresa' in linha.keys() else ""

        apo = "Sim" if str(linha['aposentado']) in ("1", "Sim", "True") else "Não"
        dom = "Sim" if 'domestica' in linha.keys() and str(linha['domestica']) in ("1", "Sim", "True") else "Não"
        tags_pessoais = f"Aposentado(a): {apo}\nDoméstica: {dom}"
        
        retornou_str = "Sim" if linha['retornou'] else "Não"

        # CÁLCULO FÉRIAS EXCEL
        status_ferias = "Desativado"
        if 'controlar_ferias' in linha.keys() and linha['controlar_ferias'] == 1 and linha['per_aquisitivo_de']:
            dias_ferias, _ = calcular_dias_beneficio_aquisitivo(cpf_func, cod_emp, cod_func, linha['per_aquisitivo_de'], linha['per_aquisitivo_ate'], linha['id'])
            status_ferias = f"{dias_ferias}/180 dias"
            if dias_ferias >= 180: status_ferias += "\n(DIREITO PERDIDO)"

        cursor.execute("SELECT codigo, descricao FROM cids_historico WHERE id_afastamento=? ORDER BY id ASC", (linha['id'],))
        cids = cursor.fetchall()
        if cids:
            cids_str = "\n".join([f"• {c['codigo']} - {c['descricao'] or 'Sem descrição'}" for c in cids])
        else:
            cids_str = linha['cid'] or ""

        cursor.execute("SELECT id, data_analise, gerou_pericia, numero_beneficio, situacao, concessao_de, concessao_ate FROM analises_inss WHERE id_afastamento=? ORDER BY id ASC", (linha['id'],))
        analises = cursor.fetchall()

        cursor.execute("SELECT data_agendamento, data_atendimento, especie, numero_beneficio, situacao, concessao_de, concessao_ate, id_analise_mae FROM pericias_inss WHERE id_afastamento=? ORDER BY id ASC", (linha['id'],))
        todas_pericias = cursor.fetchall()

        inss_parts = []
        for i, an in enumerate(analises, 1):
            de_ate = f"{an['concessao_de'] or '-'} a {an['concessao_ate'] or '-'}" if an['concessao_de'] or an['concessao_ate'] else "-"
            part = f"[{i}ª Análise Documental]\nData: {an['data_analise'] or '-'} | Situação: {an['situacao'] or '-'}\nBen: {an['numero_beneficio'] or '-'} | Período: {de_ate}"
            
            pericias_filhas = [p for p in todas_pericias if p['id_analise_mae'] == an['id']]
            for j, p in enumerate(pericias_filhas, 1):
                p_de_ate = f"{p['concessao_de'] or '-'} a {p['concessao_ate'] or '-'}" if p['concessao_de'] or p['concessao_ate'] else "-"
                part += f"\n   ↳ Perícia Vinculada {j}:\n      Agend: {p['data_agendamento'] or '-'}  Atend: {p['data_atendimento'] or '-'}\n      Sit: {p['situacao'] or '-'} | Esp: {p['especie'] or '-'}\n      Período: {p_de_ate}"
            
            inss_parts.append(part)

        pericias_orfas = [p for p in todas_pericias if not p['id_analise_mae']]
        for k, p in enumerate(pericias_orfas, 1):
            p_de_ate = f"{p['concessao_de'] or '-'} a {p['concessao_ate'] or '-'}" if p['concessao_de'] or p['concessao_ate'] else "-"
            part = f"[Perícia Sem Vínculo {k}]\nAgend: {p['data_agendamento'] or '-'}  Atend: {p['data_atendimento'] or '-'}\nSit: {p['situacao'] or '-'} | Esp: {p['especie'] or '-'}\nBen: {p['numero_beneficio'] or '-'} | Período: {p_de_ate}"
            inss_parts.append(part)

        inss_str = "\n\n".join(inss_parts) if inss_parts else "Nenhum trâmite registrado."

        cursor.execute("SELECT data_de, data_ate, qtd_dias, observacoes FROM atestados_historico WHERE id_afastamento=? ORDER BY id ASC", (linha['id'],))
        atestados = cursor.fetchall()
        atest_parts = []
        if atestados:
            for i, at in enumerate(atestados, 1):
                dias = at['qtd_dias'] or '-'
                obs_at = at['observacoes'] or 'Sem observações'
                atest_parts.append(f"[{i}º Atestado]\nPeríodo: {at['data_de'] or '-'} a {at['data_ate'] or '-'} ({dias} dias)\nObs: {obs_at}")
            atest_str = "\n\n".join(atest_parts)
        else:
            atest_str = "Nenhum atestado registrado."

        obs_gerais = linha['observacoes'] if 'observacoes' in linha.keys() and str(linha['observacoes']).strip() != "" else ""
        
        if linha['tipo_afastamento'] == "Maternidade":
            obs_mat = []
            try:
                if linha['aborto'] == 1: obs_mat.append("• Afastamento por Aborto: Sim")
                if linha['data_nascimento']: obs_mat.append(f"• Data Nascimento: {linha['data_nascimento']}")
                if linha['salario_familia']: obs_mat.append(f"• Salário Família: {linha['salario_familia']} ({linha['salario_familia_obs']})")
                if linha['imposto_renda']: obs_mat.append(f"• Imposto Renda: {linha['imposto_renda']} ({linha['imposto_renda_obs']})")
            except IndexError:
                pass
            
            if obs_mat:
                if obs_gerais: obs_gerais += "\n\n"
                obs_gerais += "[DADOS DA MATERNIDADE]\n" + "\n".join(obs_mat)

        dados_linha = [
            linha['id'], linha['empresa'] or "", cod_emp,
            linha['nome_empregado'] or "", cpf_func, cod_func,
            linha['tipo_afastamento'] or "", tags_pessoais,
            linha['ultimo_dia_trab'] or "", linha['data_obrigacao'] or "",
            linha['periodo_de'] or "", linha['periodo_ate'] or "",
            linha['qtd_dias'] or "", retornou_str, status_ferias,
            cids_str, inss_str, atest_str, obs_gerais
        ]
        
        ws.append(dados_linha)

        for cell in ws[ws.max_row]:
            cell.border = borda_fina
            if cell.column <= 15: cell.alignment = alinhamento_centro
            else: cell.alignment = alinhamento_esq_topo

    conn.close()
    auto_ajustar_colunas(ws)

    try:
        wb.save(caminho_arquivo)
        if platform.system() == 'Windows': os.startfile(caminho_arquivo)
        elif platform.system() == 'Darwin': subprocess.call(('open', caminho_arquivo))
        else: subprocess.call(('xdg-open', caminho_arquivo))
    except Exception as e:
        messagebox.showerror("Erro", f"Erro ao gerar o Excel:\n{str(e)}")

def gerar_excel_gestantes(ids_lista, caminho_arquivo):
    if not HAS_OPENPYXL:
        messagebox.showerror("Biblioteca Ausente", "Para gerar relatórios em Excel, é necessário instalar a biblioteca openpyxl.\n\nAbra o terminal do seu projeto e digite:\n\npip install openpyxl")
        return
        
    if not ids_lista:
        messagebox.showwarning("Aviso", "Não há dados na tabela para exportar.")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Acompanhamento Gestantes"

    preenchimento_cabecalho = PatternFill(start_color="DB2777", end_color="DB2777", fill_type="solid") 
    fonte_cabecalho = Font(color="FFFFFF", bold=True)
    alinhamento_centro = Alignment(horizontal="center", vertical="center")
    alinhamento_esq = Alignment(horizontal="left", vertical="center")
    borda_fina = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    cabecalhos = [
        "ID", "Empresa", "Cód Emp.", "Funcionária", "Cód Func.", 
        "Status Atual", "Data do Exame", "Semanas Const.", "Prev. Parto (40 Sem.)", "Fim da Estabilidade"
    ]
    
    ws.append(cabecalhos)

    for cell in ws[1]:
        cell.fill = preenchimento_cabecalho
        cell.font = fonte_cabecalho
        cell.alignment = alinhamento_centro
        cell.border = borda_fina

    conn = conectar_db()
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    placeholders = ','.join('?' * len(ids_lista))
    
    query = f"SELECT * FROM gestantes WHERE id IN ({placeholders}) ORDER BY empresa, nome_empregado"
    cursor.execute(query, ids_lista)
    registros = cursor.fetchall()

    for linha in registros:
        dados_linha = [
            linha['id'], linha['empresa'] or "", linha['cod_empresa'] if 'cod_empresa' in linha.keys() else "",
            linha['nome_empregado'] or "", linha['cod_empregado'] if 'cod_empregado' in linha.keys() else "",
            linha['status'] or "", linha['data_exame'] or "", linha['semanas_exame'] or "",
            linha['data_37_semanas'] or "", linha['estabilidade'] or ""
        ]
        
        ws.append(dados_linha)

        for cell in ws[ws.max_row]:
            cell.border = borda_fina
            if cell.column in [1, 3, 5, 6, 7, 8, 9, 10]: cell.alignment = alinhamento_centro
            else: cell.alignment = alinhamento_esq

    conn.close()
    auto_ajustar_colunas(ws)

    try:
        wb.save(caminho_arquivo)
        if platform.system() == 'Windows': os.startfile(caminho_arquivo)
        elif platform.system() == 'Darwin': subprocess.call(('open', caminho_arquivo))
        else: subprocess.call(('xdg-open', caminho_arquivo))
    except Exception as e:
        messagebox.showerror("Erro", f"Erro ao gerar o Excel:\n{str(e)}")